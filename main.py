from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import pandas as pd
import tempfile
import os
from urllib.parse import urlparse, parse_qs
import requests
from docx import Document

app = Flask(__name__)
CORS(app)

def extract_sheet_id(sheet_url):
    parts = sheet_url.split("/")
    return parts[5] if len(parts) > 5 else None

def download_csv(sheet_url):
    sheet_id = extract_sheet_id(sheet_url)
    if sheet_id:
        export_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv"
        response = requests.get(export_url)
        if response.status_code == 200:
            return pd.read_csv(pd.compat.StringIO(response.text))
    return pd.DataFrame()

@app.route('/get-columns', methods=['POST'])
def get_columns():
    data = request.json
    df1 = download_csv(data['sheet1'])
    df2 = download_csv(data['sheet2'])
    common_columns = list(set(df1.columns).intersection(set(df2.columns)))
    return jsonify({"columns": common_columns})

@app.route('/merge', methods=['POST'])
def merge_sheets():
    data = request.json
    df1 = download_csv(data['sheet1'])
    df2 = download_csv(data['sheet2'])
    merge_column = data['merge_column']
    selected_columns = data['selected_columns']

    merged = pd.merge(df1, df2, on=merge_column, how='inner')
    merged = merged[selected_columns]

    temp_dir = tempfile.mkdtemp()

    excel_path = os.path.join(temp_dir, 'merged.xlsx')
    csv_path = os.path.join(temp_dir, 'merged.csv')
    word_path = os.path.join(temp_dir, 'merged.docx')

    merged.to_excel(excel_path, index=False)
    merged.to_csv(csv_path, index=False)

    # Word file
    doc = Document()
    doc.add_heading('Merged Data', 0)
    table = doc.add_table(rows=1, cols=len(merged.columns))
    hdr_cells = table.rows[0].cells
    for i, col in enumerate(merged.columns):
        hdr_cells[i].text = col
    for row in merged.itertuples(index=False):
        row_cells = table.add_row().cells
        for i, value in enumerate(row):
            row_cells[i].text = str(value)
    doc.save(word_path)

    return jsonify({
        'excel': excel_path,
        'csv': csv_path,
        'word': word_path
    })

@app.route('/download', methods=['GET'])
def download():
    file_path = request.args.get('file')
    return send_file(file_path, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)